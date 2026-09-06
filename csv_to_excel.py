import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import ScatterChart, Reference, Series

CSV_PATH = 'data.csv'
XLSX_PATH = 'data.xlsx'

HEADER_FILL = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF')
TITLE_FONT = Font(bold=True, size=13)
LABEL_FONT = Font(bold=True)

# Đọc toàn bộ dữ liệu dạng chuỗi (dtype=str) để giữ nguyên các lỗi định dạng
# (khoảng trắng thừa, "2.5tr", dấu phẩy thập phân...) thay vì để pandas suy
# luận kiểu cột và "nuốt" mất lỗi khi cả cột lẫn giá trị số lẫn text.
df = pd.read_csv(CSV_PATH, dtype=str, keep_default_na=True)
n = len(df)

NUMERIC_COLS = {'Khoang_Cach_km', 'Gia_Thue_Trieu', 'So_Nguoi'}


def cell_value(col_name, raw):
    """Ô trống -> None. Ở cột số: chuỗi số hợp lệ -> number thật (để công thức
    Excel tính được); chuỗi lỗi định dạng (vd '2.5tr', '2,5') -> giữ nguyên
    text để công thức tự bỏ qua, đúng như một ô lỗi thật sự."""
    if pd.isna(raw):
        return None
    if col_name in NUMERIC_COLS:
        try:
            f = float(raw)
            return int(f) if f.is_integer() else f
        except ValueError:
            return raw
    return raw


wb = Workbook()

# ---------------------------------------------------------------
# Sheet 1: Data - dữ liệu gốc từ CSV
# ---------------------------------------------------------------
ws_data = wb.active
ws_data.title = 'Data'

ws_data.append(list(df.columns))
for _, row in df.iterrows():
    ws_data.append([cell_value(col, row[col]) for col in df.columns])

for col_idx, col_name in enumerate(df.columns, start=1):
    cell = ws_data.cell(row=1, column=col_idx)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='center')
    col_lens = df[col_name].map(lambda v: 0 if pd.isna(v) else len(str(v)))
    max_len = max(col_lens.max(), len(col_name)) + 2
    ws_data.column_dimensions[get_column_letter(col_idx)].width = min(max_len, 30)

ws_data.freeze_panes = 'A2'

# Cột tham chiếu để các công thức ở sheet khác dùng
DIST_COL = df.columns.get_loc('Khoang_Cach_km') + 1   # G
PRICE_COL = df.columns.get_loc('Gia_Thue_Trieu') + 1  # H
ROOM_COL = df.columns.get_loc('Loai_Phong') + 1        # F
AREA_COL = df.columns.get_loc('Khu_Vuc') + 1            # E
VEHICLE_COL = df.columns.get_loc('Phuong_Tien') + 1     # J
MAJOR_COL = df.columns.get_loc('Nganh') + 1              # D

DIST_L = get_column_letter(DIST_COL)
PRICE_L = get_column_letter(PRICE_COL)
ROOM_L = get_column_letter(ROOM_COL)
AREA_L = get_column_letter(AREA_COL)
VEHICLE_L = get_column_letter(VEHICLE_COL)
MAJOR_L = get_column_letter(MAJOR_COL)

DIST_RANGE = f"Data!${DIST_L}$2:${DIST_L}${n+1}"
PRICE_RANGE = f"Data!${PRICE_L}$2:${PRICE_L}${n+1}"
ROOM_RANGE = f"Data!${ROOM_L}$2:${ROOM_L}${n+1}"
AREA_RANGE = f"Data!${AREA_L}$2:${AREA_L}${n+1}"
VEHICLE_RANGE = f"Data!${VEHICLE_L}$2:${VEHICLE_L}${n+1}"
MAJOR_RANGE = f"Data!${MAJOR_L}$2:${MAJOR_L}${n+1}"

# ---------------------------------------------------------------
# Sheet 2: ThongKe - thống kê mô tả (công thức Excel sống)
# ---------------------------------------------------------------
ws_stat = wb.create_sheet('ThongKe')
ws_stat['A1'] = 'THỐNG KÊ MÔ TẢ'
ws_stat['A1'].font = TITLE_FONT

headers = ['Chỉ số', 'Khoảng cách (km)', 'Giá thuê (triệu)']
for i, h in enumerate(headers, start=1):
    c = ws_stat.cell(row=3, column=i, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL

rows = [
    ('Số quan sát (N)', f'=COUNT({DIST_RANGE})', f'=COUNT({PRICE_RANGE})'),
    ('Trung bình (Mean)', f'=AVERAGE({DIST_RANGE})', f'=AVERAGE({PRICE_RANGE})'),
    ('Trung vị (Median)', f'=MEDIAN({DIST_RANGE})', f'=MEDIAN({PRICE_RANGE})'),
    ('Độ lệch chuẩn (StDev)', f'=STDEV({DIST_RANGE})', f'=STDEV({PRICE_RANGE})'),
    ('Phương sai (Variance)', f'=VAR({DIST_RANGE})', f'=VAR({PRICE_RANGE})'),
    ('Nhỏ nhất (Min)', f'=MIN({DIST_RANGE})', f'=MIN({PRICE_RANGE})'),
    ('Lớn nhất (Max)', f'=MAX({DIST_RANGE})', f'=MAX({PRICE_RANGE})'),
    ('Khoảng biến thiên (Range)', f'=MAX({DIST_RANGE})-MIN({DIST_RANGE})', f'=MAX({PRICE_RANGE})-MIN({PRICE_RANGE})'),
]
for r_offset, (label, f_dist, f_price) in enumerate(rows, start=4):
    ws_stat.cell(row=r_offset, column=1, value=label).font = LABEL_FONT
    ws_stat.cell(row=r_offset, column=2, value=f_dist).number_format = '0.0000'
    ws_stat.cell(row=r_offset, column=3, value=f_price).number_format = '0.0000'

ws_stat.column_dimensions['A'].width = 26
ws_stat.column_dimensions['B'].width = 20
ws_stat.column_dimensions['C'].width = 20

corr_row = 4 + len(rows) + 1
ws_stat.cell(row=corr_row, column=1, value='Hệ số tương quan (Correlation)').font = LABEL_FONT
ws_stat.cell(row=corr_row, column=2, value=f'=CORREL({DIST_RANGE},{PRICE_RANGE})').number_format = '0.0000'

# ---------------------------------------------------------------
# Sheet 3: HoiQuy - hồi quy tuyến tính đơn biến (công thức Excel)
# ---------------------------------------------------------------
ws_reg = wb.create_sheet('HoiQuy')
ws_reg['A1'] = 'HỒI QUY TUYẾN TÍNH ĐƠN BIẾN: Giá thuê ~ Khoảng cách'
ws_reg['A1'].font = TITLE_FONT

ws_reg['A3'] = 'Hệ số góc (Slope, β1)'
ws_reg['B3'] = f'=SLOPE({PRICE_RANGE},{DIST_RANGE})'
ws_reg['A4'] = 'Hệ số chặn (Intercept, β0)'
ws_reg['B4'] = f'=INTERCEPT({PRICE_RANGE},{DIST_RANGE})'
ws_reg['A5'] = 'Hệ số xác định (R-squared)'
ws_reg['B5'] = f'=RSQ({PRICE_RANGE},{DIST_RANGE})'
ws_reg['A6'] = 'Hệ số tương quan (R)'
ws_reg['B6'] = f'=CORREL({DIST_RANGE},{PRICE_RANGE})'
ws_reg['A7'] = 'Sai số chuẩn của ước lượng (StEyx)'
ws_reg['B7'] = f'=STEYX({PRICE_RANGE},{DIST_RANGE})'

for r in range(3, 8):
    ws_reg.cell(row=r, column=1).font = LABEL_FONT
    ws_reg.cell(row=r, column=2).number_format = '0.0000'

ws_reg['A9'] = 'Phương trình hồi quy:'
ws_reg['A9'].font = LABEL_FONT
ws_reg['A10'] = '=CONCATENATE("Giá thuê = ", TEXT(B4,"0.00"), " + (", TEXT(B3,"0.00"), ") * Khoảng cách")'
ws_reg.column_dimensions['A'].width = 30
ws_reg.column_dimensions['B'].width = 16

# Bảng dự đoán & phần dư (residuals) cho toàn bộ dữ liệu
ws_reg['D2'] = 'STT'
ws_reg['E2'] = 'Khoảng cách (km)'
ws_reg['F2'] = 'Giá thực tế'
ws_reg['G2'] = 'Giá dự đoán'
ws_reg['H2'] = 'Phần dư (Thực tế - Dự đoán)'
for col in ['D', 'E', 'F', 'G', 'H']:
    ws_reg[f'{col}2'].font = HEADER_FONT
    ws_reg[f'{col}2'].fill = HEADER_FILL
    ws_reg.column_dimensions[col].width = 16

for i in range(n):
    r = i + 3
    data_row = i + 2
    ws_reg.cell(row=r, column=4, value=i + 1)
    ws_reg.cell(row=r, column=5, value=f'=Data!{DIST_L}{data_row}')
    ws_reg.cell(row=r, column=6, value=f'=Data!{PRICE_L}{data_row}')
    ws_reg.cell(row=r, column=7, value=f'=$B$4+$B$3*E{r}').number_format = '0.0000'
    ws_reg.cell(row=r, column=8, value=f'=F{r}-G{r}').number_format = '0.0000'

# Biểu đồ phân tán + đường hồi quy
chart = ScatterChart()
chart.title = 'Khoảng cách vs Giá thuê (kèm giá trị dự đoán)'
chart.x_axis.title = 'Khoảng cách (km)'
chart.y_axis.title = 'Giá thuê (triệu)'
chart.height = 10
chart.width = 18

x_values = Reference(ws_reg, min_col=5, min_row=3, max_row=n + 2)
y_actual = Reference(ws_reg, min_col=6, min_row=2, max_row=n + 2)
y_pred = Reference(ws_reg, min_col=7, min_row=2, max_row=n + 2)

s1 = Series(y_actual, x_values, title_from_data=True)
s1.marker.symbol = 'circle'
s1.graphicalProperties.line.noFill = True
chart.series.append(s1)

s2 = Series(y_pred, x_values, title_from_data=True)
s2.marker.symbol = 'none'
chart.series.append(s2)

ws_reg.add_chart(chart, 'J2')

# ---------------------------------------------------------------
# Sheet 4: PhanTichNhom - phân tích theo nhóm (loại phòng, khu vực, phương tiện, ngành)
# ---------------------------------------------------------------
ws_grp = wb.create_sheet('PhanTichNhom')
ws_grp['A1'] = 'PHÂN TÍCH GIÁ THUÊ THEO NHÓM'
ws_grp['A1'].font = TITLE_FONT

def write_group_table(ws, start_row, title, group_range, categories):
    ws.cell(row=start_row, column=1, value=title).font = LABEL_FONT
    hdr = ['Nhóm', 'Số lượng', 'Giá TB (triệu)', 'Khoảng cách TB (km)']
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=start_row + 1, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    r = start_row + 2
    for cat in categories:
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=f'=COUNTIF({group_range},A{r})')
        ws.cell(row=r, column=3, value=f'=AVERAGEIF({group_range},A{r},{PRICE_RANGE})').number_format = '0.0000'
        ws.cell(row=r, column=4, value=f'=AVERAGEIF({group_range},A{r},{DIST_RANGE})').number_format = '0.0000'
        r += 1
    return r + 1

next_row = write_group_table(ws_grp, 3, 'Theo Loại phòng', ROOM_RANGE, sorted(df['Loai_Phong'].dropna().unique()))
next_row = write_group_table(ws_grp, next_row, 'Theo Khu vực', AREA_RANGE, sorted(df['Khu_Vuc'].dropna().unique()))
next_row = write_group_table(ws_grp, next_row, 'Theo Phương tiện di chuyển', VEHICLE_RANGE, sorted(df['Phuong_Tien'].dropna().unique()))
write_group_table(ws_grp, next_row, 'Theo Ngành học', MAJOR_RANGE, sorted(df['Nganh'].dropna().unique()))

for col, w in zip(['A', 'B', 'C', 'D'], [26, 12, 16, 20]):
    ws_grp.column_dimensions[col].width = w

wb.save(XLSX_PATH)
print(f"Đã tạo {XLSX_PATH} với {n} dòng dữ liệu và các sheet: {wb.sheetnames}")
